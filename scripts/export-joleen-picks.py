#!/usr/bin/env python3
"""
从 YoHo 创意榜单 xlsx + Desktop/joleen 导出前三类 Featured 精选素材。

默认合并两份榜单（先 2026 再 2025，按表内行序、同项目夹去重）：
  ranklist-YoHo-2026-01-01-2026-09-23-....xlsx
  ranklist-YoHo-2025-01-01-2025-12-31-....xlsx

用法：
  python3 scripts/export-joleen-picks.py
  python3 scripts/export-joleen-picks.py /path/a.xlsx /path/b.xlsx
"""
import json
import re
import shutil
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

try:
    import openpyxl
except ImportError:
    raise SystemExit("请先安装：pip3 install openpyxl")

ROOT = Path(__file__).resolve().parents[1]
JOLEEN = Path("/Users/mico/Desktop/joleen")
OUT = Path("/Users/mico/Desktop/Joleen-精选素材")
DEFAULT_XLSX = Path(
    "/Users/mico/Downloads/ranklist-YoHo-2026-01-01-2026-09-23-2026-09-24-18-54-47-ne.xlsx"
)
RANKLIST_2025 = Path(
    "/Users/mico/Downloads/ranklist-YoHo-2025-01-01-2025-12-31-2026-09-24-19-18-58-dj.xlsx"
)

CAT_DIRS = {
    "feed-video": OUT / "01-信息流-视频",
    "feed-image": OUT / "02-信息流-图片",
    "aigc": OUT / "03-AIGC",
}

COL = {
    "name": 1,
    "preview": 2,
    "video_url": 3,
    "topic": 15,
    "region": 16,
    "source": 17,
    "spend": 23,
    "roi7": 27,
    "type": 6,
    "channel": 7,
    "cost": 31,
}

FEED_VIDEO_LIMIT = 15
FEED_IMAGE_LIMIT = 12
AIGC_LIMIT = 10

REGION_ZONE = {
    "台湾": ("台湾", "华语", "中国台湾"),
    "土耳其": ("土耳其", "土耳其区"),
    "中东": ("中东", "伊拉克", "沙特", "叙利亚", "埃及", "阿曼", "也门", "科威特", "卡塔尔", "阿联酋", "AR"),
}


def parse_money(v: Any) -> float:
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").strip()
    try:
        return float(re.sub(r"[^\d.]", "", s) or 0)
    except ValueError:
        return 0.0


def parse_roi(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        return f"{v:.2f}%"
    s = str(v).strip()
    return s if "%" in s else f"{s}%"


def parse_roi_num(v: Any) -> float:
    return parse_money(str(v).replace("%", "") if v is not None else "")


def market_zone(region: Any, creative_name: str) -> str:
    text = f"{region or ''} {creative_name}"
    for zone, keys in REGION_ZONE.items():
        if any(k in text for k in keys):
            return zone
    m = re.search(r"YOHO-([A-Z]{2})", creative_name, re.I)
    if m:
        code = m.group(1).upper()
        if code == "TW":
            return "台湾"
        if code == "TR":
            return "土耳其"
        if code in ("AR", "EG"):
            return "中东"
    return str(region or "—")[:12]


def period_from_name(name: str) -> str:
    m = re.search(r"(\d{8})", name)
    return f"{m.group(1)}期" if m else "—"


def channel_label(v: Any) -> str:
    s = str(v or "").strip()
    return s if s else "Facebook"


def classify_row(name: str, creative_type: Any) -> Optional[str]:
    ctype = str(creative_type or "").upper()
    if "街访" in name:
        return "aigc"
    if ctype == "IMAGE" or "图片" in name:
        return "feed-image"
    if ctype == "VIDEO" or "视频" in name:
        return "feed-video"
    return None


def folder_key_from_creative(name: str) -> str:
    """从创意文件名提取与 joleen 项目夹对齐的键（日期+产品+主题段）。"""
    base = Path(name).name
    base = re.sub(r"\.(mp4|mov|jpg|jpeg|png)$", "", base, flags=re.I)
    base = re.sub(r"_edit_\d+$", "", base, flags=re.I)
    base = re.sub(r"^\d+-\d+-", "", base)
    base = re.sub(r"^\d+-\d+-\d+-", "", base)
    return base


def build_joleen_index() -> Dict[str, Path]:
    index: Dict[str, Path] = {}
    for d in JOLEEN.iterdir():
        if not d.is_dir():
            continue
        norm = d.name
        index[norm.lower()] = d
        # 无 # 的旧命名
        alt = norm.replace("#", "-")
        index[alt.lower()] = d
    return index


def find_joleen_folder(creative_name: str, index: Dict[str, Path]) -> Optional[Path]:
    key = folder_key_from_creative(creative_name)
    kl = key.lower()
    if kl in index:
        return index[kl]
    # 前缀匹配：20260518-yoho-tr#双女口播
    for name, path in index.items():
        if kl in name or name in kl:
            return path
    m = re.search(r"(\d{8})-([^#]+#[^#]+)", key, re.I)
    if m:
        date, rest = m.group(1), m.group(2).lower()
        for name, path in index.items():
            if date in name and rest.split("#")[0] in name:
                return path
    return None


def load_ranklist(xlsx: Path) -> List[dict]:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb.active
    rows: List[dict] = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, COL["name"]).value
        if not name:
            continue
        name = str(name).strip()
        cat = classify_row(name, ws.cell(r, COL["type"]).value)
        if not cat:
            continue
        preview = ws.cell(r, COL["preview"]).value or ""
        video_url = ws.cell(r, COL["video_url"]).value or ""
        rows.append(
            {
                "rank": r - 1,
                "name": name,
                "category": cat,
                "topic": str(ws.cell(r, COL["topic"]).value or "").strip(),
                "region": ws.cell(r, COL["region"]).value or "",
                "source": ws.cell(r, COL["source"]).value or "",
                "spend": parse_money(ws.cell(r, COL["spend"]).value),
                "roi7": parse_roi(ws.cell(r, COL["roi7"]).value),
                "roi7_num": parse_roi_num(ws.cell(r, COL["roi7"]).value),
                "video_url": video_url,
                "preview": preview,
                "cost": parse_money(ws.cell(r, COL["cost"]).value),
                "channel": channel_label(ws.cell(r, COL["channel"]).value),
                "marketZone": market_zone(
                    ws.cell(r, COL["region"]).value, name
                ),
                "period": period_from_name(name),
            }
        )
    wb.close()
    return rows


def pick_by_rank(
    rows: List[dict], category: str, index: Dict[str, Path], limit: int
) -> List[dict]:
    seen: set[str] = set()
    picks: List[dict] = []
    for row in rows:
        if row["category"] != category:
            continue
        folder = find_joleen_folder(row["name"], index)
        if not folder:
            continue
        fid = str(folder.resolve())
        if fid in seen:
            continue
        seen.add(fid)
        picks.append({**row, "folder": folder})
        if len(picks) >= limit:
            break
    return picks


def merge_ranklists(paths: List[Path]) -> List[dict]:
    merged: List[dict] = []
    for xlsx in paths:
        if not xlsx.is_file():
            raise SystemExit(f"找不到榜单文件：{xlsx}")
        merged.extend(load_ranklist(xlsx))
    return merged


def cover_from_pick(pick: dict) -> str:
    preview = str(pick.get("preview") or "").strip()
    if preview:
        return preview
    video = str(pick.get("video_url") or "").split("?")[0]
    return video


def main() -> None:
    xlsx_paths = [Path(a) for a in sys.argv[1:]] if len(sys.argv) > 1 else [DEFAULT_XLSX, RANKLIST_2025]
    if not JOLEEN.is_dir():
        raise SystemExit(f"找不到 joleen 文件夹：{JOLEEN}")

    rows = merge_ranklists(xlsx_paths)
    index = build_joleen_index()

    if OUT.exists():
        shutil.rmtree(OUT)
    for p in CAT_DIRS.values():
        p.mkdir(parents=True)

    manifest: List[Tuple[str, List[str]]] = []

    def copy_picks(cat: str, limit: int) -> None:
        picks = pick_by_rank(rows, cat, index, limit)
        lines: List[str] = []
        dest_root = CAT_DIRS[cat]
        for i, pick in enumerate(picks, 1):
            src = pick["folder"]
            dest = dest_root / f"{i:02d}-{src.name}"
            shutil.copytree(src, dest)
            title = pick["topic"] or pick["name"]
            lines.append(
                f"{i:02d}. 榜单#{pick['rank']} {title}\n"
                f"    spend {pick['spend']:.2f} · 7日ROI {pick['roi7']} · {pick['region']}\n"
                f"    {src.name}"
            )
        manifest.append((dest_root.name, lines))

    copy_picks("feed-video", FEED_VIDEO_LIMIT)
    copy_picks("feed-image", FEED_IMAGE_LIMIT)
    copy_picks("aigc", AIGC_LIMIT)

    readme = OUT / "README-筛选说明.txt"
    readme.write_text(
        "Joleen 精选素材（数据源：YoHo 创意榜单 xlsx）\n"
        f"榜单文件：\n"
        + "\n".join(f"  - {p}" for p in xlsx_paths)
        + "\n"
        "分类：街访→AIGC；IMAGE/图片→信息流·图片；其余 VIDEO→信息流·视频\n"
        "排序：按榜单行顺序（自上而下为优先级），同一项目夹只保留一次\n\n"
        + "\n".join(f"【{name}】\n" + "\n".join(lines) + "\n" for name, lines in manifest),
        encoding="utf-8",
    )

    # 供网站后续同步用的 JSON（不自动改 moreWork.ts）
    limits = {
        "feed-video": FEED_VIDEO_LIMIT,
        "feed-image": FEED_IMAGE_LIMIT,
        "aigc": AIGC_LIMIT,
    }
    categories = {
        cat: pick_by_rank(rows, cat, index, limits[cat]) for cat in limits
    }
    summary = {
        "sourceXlsx": [str(p) for p in xlsx_paths],
        "exportedAt": __import__("datetime").datetime.now().isoformat(timespec="seconds"),
        "categories": {
            cat: [
                {
                    "rank": p["rank"],
                    "topic": p["topic"],
                    "spend": p["spend"],
                    "roi7": p["roi7"],
                    "cost": p.get("cost", 0),
                    "marketZone": p.get("marketZone"),
                    "channel": p.get("channel"),
                    "source": p.get("source"),
                    "period": p.get("period"),
                    "folder": p["folder"].name,
                    "creativeName": p["name"],
                    "preview": str(p.get("preview") or ""),
                    "cover": cover_from_pick(p),
                    "videoUrl": str(p["video_url"]).split("?")[0] if p["video_url"] else "",
                }
                for p in categories[cat]
            ]
            for cat in limits
        },
    }
    (ROOT / "data" / "ranklist-picks.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print(f"Done: {OUT}")
    print(f"Summary: {ROOT / 'data' / 'ranklist-picks.json'}")


if __name__ == "__main__":
    main()
